import asyncio
import io
import re

import edge_tts
import pypdfium2 as pdfium
import pytesseract
import streamlit as st
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Path to Tesseract (change this if you installed it somewhere else)
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

st.set_page_config(page_title="Document to Voice", page_icon="🎧")
st.title("🎧 Text / PDF / PPT → Voice (MP3)")
st.caption("Type or paste text, or upload a PDF / PowerPoint file. Choose a voice and download the MP3.")

VOICES = {
    "English (US)": {"Female": "en-US-JennyNeural", "Male": "en-US-GuyNeural"},
    "English (UK)": {"Female": "en-GB-SoniaNeural", "Male": "en-GB-RyanNeural"},
    "Korean": {"Female": "ko-KR-SunHiNeural", "Male": "ko-KR-InJoonNeural"},
    "Spanish": {"Female": "es-ES-ElviraNeural", "Male": "es-ES-AlvaroNeural"},
    "French": {"Female": "fr-FR-DeniseNeural", "Male": "fr-FR-HenriNeural"},
    "Urdu": {"Female": "ur-PK-UzmaNeural", "Male": "ur-PK-AsadNeural"},
    "Hindi": {"Female": "hi-IN-SwaraNeural", "Male": "hi-IN-MadhurNeural"},
}
OCR_LANG = {
    "English (US)": "eng",
    "English (UK)": "eng",
    "Korean": "kor+eng",
    "Spanish": "spa+eng",
    "French": "fra+eng",
    "Urdu": "urd+eng",
    "Hindi": "hin+eng",
}


# ---------- OCR ----------
def ocr_image(img: Image.Image, lang: str) -> str:
    return pytesseract.image_to_string(img, lang=lang)


def ocr_pdf(data: bytes, lang: str) -> str:
    pdf = pdfium.PdfDocument(data)
    total = len(pdf)
    bar = st.progress(0, text="Running OCR (this may take a while)...")
    texts = []
    for i in range(total):
        img = pdf[i].render(scale=200 / 72).to_pil()
        texts.append(ocr_image(img, lang))
        bar.progress((i + 1) / total, text=f"OCR: page {i + 1} of {total}")
    bar.empty()
    return "\n".join(texts)


# ---------- Text extraction ----------
def extract_pdf(data: bytes, lang: str) -> str:
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    if len(text.strip()) < 20:  # no real text found -> use OCR
        text = ocr_pdf(data, lang)
    return text


def _shape_text(shape, pics):
    out = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            out.extend(_shape_text(s, pics))
    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        pics.append(shape.image.blob)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                out.append(cell.text)
    elif shape.has_text_frame:
        out.append(shape.text_frame.text)
    return out


def extract_pptx(data: bytes, lang: str) -> str:
    prs = Presentation(io.BytesIO(data))
    out, pics = [], []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"Slide {i}.")
        for shape in slide.shapes:
            out.extend(_shape_text(shape, pics))
        if slide.has_notes_slide:
            out.append(slide.notes_slide.notes_text_frame.text)
    text = "\n".join(out)
    # If there is no text, run OCR on the pictures inside the slides
    if len(re.sub(r"Slide \d+\.", "", text).strip()) < 20 and pics:
        bar = st.progress(0, text="Running OCR on slide images...")
        parts = []
        for i, blob in enumerate(pics):
            parts.append(ocr_image(Image.open(io.BytesIO(blob)), lang))
            bar.progress((i + 1) / len(pics))
        bar.empty()
        text = "\n".join(parts)
    return text


@st.cache_data(show_spinner=False)
def get_text(data: bytes, name: str, lang: str) -> str:
    if name.lower().endswith(".pdf"):
        return extract_pdf(data, lang)
    return extract_pptx(data, lang)


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def chunk_text(text: str, limit: int = 3000):
    sentences = re.split(r"(?<=[.!?۔।])\s+", text)
    chunks, cur = [], ""
    for s in sentences:
        if len(cur) + len(s) > limit and cur:
            chunks.append(cur)
            cur = ""
        cur += s + " "
    if cur.strip():
        chunks.append(cur)
    return chunks


# ---------- Text to speech ----------
async def synth(text, voice, rate):
    comm = edge_tts.Communicate(text, voice, rate=rate)
    buf = b""
    async for item in comm.stream():
        if item["type"] == "audio":
            buf += item["data"]
    return buf


def text_to_mp3(text, voice, rate, progress):
    chunks = chunk_text(text)
    audio = b""
    for i, c in enumerate(chunks):
        audio += asyncio.run(synth(c, voice, rate))
        progress.progress((i + 1) / len(chunks))
    return audio


# ---------- UI ----------
mode = st.radio("Input type", ["Type / paste text", "Upload file"], horizontal=True)

col1, col2 = st.columns(2)
lang = col1.selectbox("Language (of the text and the voice)", list(VOICES.keys()))
gender = col2.radio("Voice", ["Female", "Male"], horizontal=True)
speed = st.slider("Speed", -30, 30, 0, help="0 = normal, negative = slower, positive = faster")

text = ""
file_stem = "audio"

if mode == "Type / paste text":
    text = st.text_area("Enter your text here", height=250, placeholder="Type or paste your text...")
    text = clean(text)
else:
    file = st.file_uploader("Upload a PDF or PPTX file", type=["pdf", "pptx"])
    if file:
        file_stem = file.name.rsplit(".", 1)[0]
        try:
            raw = get_text(file.getvalue(), file.name, OCR_LANG[lang])
        except pytesseract.TesseractNotFoundError:
            st.error("Tesseract is not installed or the path is wrong. Please check the installation.")
            st.stop()
        except pytesseract.TesseractError as e:
            st.error(f"OCR error (the {lang} language data may not be installed): {e}")
            st.stop()

        text = clean(raw)
        if not text:
            st.error("No text found. The image may be unclear or the wrong language is selected.")
        else:
            text = st.text_area("Extracted text (you can edit it)", text, height=250)

if text:
    st.caption(f"{len(text)} characters")
    if st.button("🔊 Generate voice", type="primary"):
        voice = VOICES[lang][gender]
        bar = st.progress(0)
        mp3 = text_to_mp3(text, voice, f"{speed:+d}%", bar)
        st.success("Done! Your audio is ready.")
        st.audio(mp3, format="audio/mp3")
        st.download_button(
            "⬇️ Download MP3",
            data=io.BytesIO(mp3),
            file_name=file_stem + ".mp3",
            mime="audio/mpeg",
        )